#!/usr/bin/env python3
"""Voice-of-customer swipe file: scrape reviews and posts into one sheet, then pull hooks from it.

  python research/voc.py collect --config research/gut-complex.toml            # every source
  python research/voc.py collect --config research/gut-complex.toml --only amazon,reddit
  python research/voc.py import research/data/claude-*.jsonl                   # rows found by hand or by Claude
  python research/voc.py hooks --theme urgency --n 30                           # hook lines from the sheet
  python research/voc.py stats

Scraping runs through Apify (one APIFY_TOKEN covers Amazon, Reddit, X/Twitter and Google).
Rows land in research/data/swipe.csv (the source of truth) and research/data/swipe.xlsx.
"""
import argparse
import csv
import hashlib
import json
import os
import re
import sys
import time
import tomllib
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import requests

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
SHEET_CSV = DATA / "swipe.csv"
SHEET_XLSX = DATA / "swipe.xlsx"

COLUMNS = [
    "id", "collected_at", "source", "platform", "kind", "product", "brand", "rating", "title", "text",
    "url", "author", "date", "community", "theme", "emotion", "pain_score", "about_product",
    "key_phrases", "verified", "query",
]

ACTORS = {
    "amazon": "junglee~amazon-reviews-scraper",
    "reddit": "trudax~reddit-scraper-lite",
    "twitter": "apidojo~tweet-scraper",
    "google": "apify~google-search-scraper",
}

# Words that signal emotional pain rather than a neutral description. Weighted by how raw they read.
PAIN_WORDS = {
    3: ["suicidal", "want to die", "rock bottom", "can't live like this", "cant live like this", "ruined my life",
        "ruining my life", "hopeless", "crying", "cried", "sobbing", "breaking down", "broke down", "desperate",
        "humiliat", "shit myself", "shat myself", "accident", "pooped my pants", "didn't make it", "didnt make it",
        "is this my life", "give up", "gave up", "no life", "prisoner", "trapped", "nightmare", "hell"],
    2: ["embarrass", "ashamed", "shame", "terrified", "scared", "afraid", "panic", "anxiety", "anxious", "dread",
        "depressed", "depression", "miserable", "exhausted", "frustrat", "furious", "angry", "gaslit", "dismissed",
        "nobody believes", "no one believes", "isolated", "alone", "lonely", "cancel", "missed", "lost my job",
        "can't leave the house", "cant leave the house", "map out", "mapping", "every bathroom", "look pregnant",
        "looks pregnant", "pregnant", "hate my body", "waste of money", "wasted", "came back", "comes back",
        "worse", "nothing works", "nothing helps", "tried everything", "years"],
    1: ["bloat", "urgency", "urgent", "rush to the bathroom", "run to the bathroom", "diarrhea", "gas", "pain",
        "cramp", "imodium", "rifaximin", "xifaxan", "probiotic", "doctor", "gastro", "gi ", "fodmap", "scared to eat",
        "afraid to eat", "every meal", "after eating", "after i eat", "relapse", "tired", "sick of"],
}
FIRST_PERSON = re.compile(r"\b(i|i'm|im|i've|ive|me|my|myself)\b", re.I)


def now():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")


def clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def row_id(url, text):
    key = (url or "") + "|" + re.sub(r"\W+", "", (text or "").lower())[:400]
    return hashlib.sha1(key.encode()).hexdigest()[:12]


def pain_score(text):
    t = (text or "").lower()
    score = sum(w for w, words in PAIN_WORDS.items() for p in words if p in t)
    if FIRST_PERSON.search(t):
        score += 2
    return score


def tag(text, mapping):
    t = (text or "").lower()
    return ", ".join(k for k, words in mapping.items() if any(w.lower() in t for w in words))


# ---------- sheet ----------

def load_sheet():
    if not SHEET_CSV.exists():
        return []
    with SHEET_CSV.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def save_sheet(rows):
    DATA.mkdir(parents=True, exist_ok=True)
    rows.sort(key=lambda r: (r.get("source", ""), -float(r.get("pain_score") or 0)))
    with SHEET_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        print("openpyxl not installed: wrote CSV only", file=sys.stderr)
        return
    wb = Workbook()
    for i, (name, keep) in enumerate([
        ("All", lambda r: True),
        ("Reviews 1-3 star", lambda r: r.get("kind") == "review"),
        ("Forums & social", lambda r: r.get("kind") in ("post", "comment", "tweet", "search")),
    ]):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append(COLUMNS)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            if keep(r):
                ws.append([r.get(c, "") for c in COLUMNS])
        ws.freeze_panes = "A2"
        widths = {"text": 90, "title": 30, "url": 40, "key_phrases": 40, "product": 22}
        for idx, col in enumerate(COLUMNS, 1):
            ws.column_dimensions[ws.cell(1, idx).column_letter].width = widths.get(col, 13)
        for row in ws.iter_rows(min_row=2):
            row[COLUMNS.index("text")].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(SHEET_XLSX)


def merge(new_rows, cfg=None):
    rows = load_sheet()
    seen = {r["id"] for r in rows}
    added = 0
    themes = (cfg or {}).get("themes", {})
    emotions = (cfg or {}).get("emotions", {})
    product_terms = [t.lower() for t in (cfg or {}).get("product_terms", [])]
    for r in new_rows:
        r["text"] = clean(r.get("text"))
        if len(r["text"]) < 25:
            continue
        r["id"] = r.get("id") or row_id(r.get("url"), r["text"])
        if r["id"] in seen:
            continue
        r.setdefault("collected_at", now())
        if not r.get("pain_score"):
            r["pain_score"] = pain_score(r["title"] + " " + r["text"] if r.get("title") else r["text"])
        for col, mapping in (("theme", themes), ("emotion", emotions)):
            tags = [t.strip() for t in (r.get(col) or "").split(",") if t.strip()]
            tags += [t for t in tag(r["text"], mapping).split(", ") if t and t not in tags]
            r[col] = ", ".join(tags)
        if product_terms and r.get("kind") != "review" and not r.get("about_product"):
            r["about_product"] = "yes" if any(p in r["text"].lower() for p in product_terms) else ""
        if isinstance(r.get("key_phrases"), list):
            r["key_phrases"] = " | ".join(r["key_phrases"])
        seen.add(r["id"])
        rows.append(r)
        added += 1
    save_sheet(rows)
    return added, len(rows)


# ---------- Apify sources ----------

def apify(actor, payload, token, timeout=600):
    url = f"https://api.apify.com/v2/acts/{actor}/run-sync-get-dataset-items"
    r = requests.post(url, params={"token": token, "timeout": timeout}, json=payload, timeout=timeout + 60)
    if r.status_code >= 400:
        raise RuntimeError(f"{actor}: HTTP {r.status_code} {r.text[:300]}")
    return r.json()


def collect_amazon(cfg, token):
    src = cfg.get("amazon", {})
    out = []
    for p in src.get("products", []):
        items = apify(ACTORS["amazon"], {
            "productUrls": [{"url": u} for u in p["urls"]],
            "maxReviews": src.get("max_reviews_per_product", 100),
            "filterByRatings": src.get("ratings", ["oneStar", "twoStar", "threeStar"]),
            "sort": src.get("sort", "helpful"),
        }, token)
        for it in items:
            rating = it.get("ratingScore")
            if it.get("foundNoReviews") or rating is None or rating > src.get("max_rating", 3):
                continue
            out.append({
                "source": "amazon", "platform": "amazon", "kind": "review",
                "product": p["name"], "brand": p.get("brand", ""), "rating": rating,
                "title": clean(it.get("reviewTitle")), "text": it.get("reviewDescription"),
                "url": it.get("reviewUrl") or (p["urls"][0] if p["urls"] else ""),
                "author": it.get("username") or "", "date": (it.get("date") or "")[:10],
                "verified": "scraped",
            })
        print(f"  amazon · {p['name']}: {len(items)} reviews fetched")
    return out


def collect_reddit(cfg, token):
    src = cfg.get("reddit", {})
    out = []
    subs = src.get("subreddits") or [""]
    for sub in subs:
        payload = {
            "searches": src["queries"], "searchPosts": True, "searchComments": src.get("comments", True),
            "sort": src.get("sort", "relevance"), "time": src.get("time", "all"),
            "maxItems": src.get("max_items_per_subreddit", 150), "maxPostCount": src.get("max_posts_per_query", 25),
            "maxComments": src.get("max_comments_per_post", 15), "skipCommunity": True, "skipUserPosts": True,
            "includeNSFW": False, "proxy": {"useApifyProxy": True, "apifyProxyGroups": ["RESIDENTIAL"]},
        }
        if sub:
            payload["searchCommunityName"] = sub
        items = apify(ACTORS["reddit"], payload, token)
        for it in items:
            kind = "comment" if it.get("dataType") == "comment" else "post"
            text = it.get("body") or ""
            out.append({
                "source": "reddit", "platform": "reddit", "kind": kind, "title": clean(it.get("title")),
                "text": text, "url": it.get("url") or "", "author": it.get("username") or "",
                "date": (it.get("createdAt") or "")[:10], "community": it.get("communityName") or "",
                "verified": "scraped",
            })
        print(f"  reddit · r/{sub or 'all'}: {len(items)} items")
    return out


def collect_twitter(cfg, token):
    src = cfg.get("twitter", {})
    items = apify(ACTORS["twitter"], {
        "searchTerms": [f"{q} -filter:links -filter:replies lang:en" if src.get("strict") else q for q in src["queries"]],
        "maxItems": src.get("max_items", 500), "sort": src.get("sort", "Top"), "tweetLanguage": "en",
    }, token)
    out = []
    for it in items:
        if it.get("isRetweet") or it.get("noResults"):
            continue
        out.append({
            "source": "twitter", "platform": "x", "kind": "tweet", "text": it.get("text") or "",
            "url": it.get("url") or it.get("twitterUrl") or "",
            "author": ((it.get("author") or {}).get("userName")) or "", "date": clean(it.get("createdAt"))[:16],
            "verified": "scraped",
        })
    print(f"  twitter: {len(items)} tweets")
    return out


def collect_google(cfg, token):
    """Google finds forum threads Reddit search misses (HealingWell, Mayo Connect, Patient.info, Quora...).
    The snippet is stored; with fetch_pages the readable paragraphs of each page are stored too."""
    src = cfg.get("google", {})
    items = apify(ACTORS["google"], {
        "queries": "\n".join(src["queries"]), "maxPagesPerQuery": src.get("pages_per_query", 1),
        "countryCode": src.get("country", "us"), "languageCode": "en",
    }, token)
    out, urls = [], []
    for it in items:
        q = (it.get("searchQuery") or {}).get("term", "")
        for res in it.get("organicResults") or []:
            url = res.get("url") or ""
            out.append({
                "source": "google", "platform": urlparse(url).netloc.replace("www.", ""), "kind": "search",
                "title": clean(res.get("title")), "text": res.get("description") or "", "url": url,
                "date": res.get("date") or "", "verified": "snippet", "query": q,
            })
            urls.append((url, q))
    print(f"  google: {len(out)} results")
    if src.get("fetch_pages"):
        out += fetch_pages(urls, src.get("skip_domains", []), src.get("min_paragraph_score", 5))
    return out


def fetch_pages(urls, skip_domains, min_score):
    from bs4 import BeautifulSoup
    out = []
    headers = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_5) AppleWebKit/537.36 Chrome/126 Safari/537.36"}
    for url, q in dict(urls).items():
        host = urlparse(url).netloc.replace("www.", "")
        if any(d in host for d in skip_domains):
            continue
        try:
            r = requests.get(url, headers=headers, timeout=20)
            if r.status_code != 200:
                continue
        except requests.RequestException:
            continue
        soup = BeautifulSoup(r.text, "html.parser")
        for p in soup.find_all(["p", "blockquote", "div"], limit=4000):
            if p.find(["p", "div"]):
                continue
            text = clean(p.get_text(" "))
            if 80 <= len(text) <= 2500 and pain_score(text) >= min_score and FIRST_PERSON.search(text):
                out.append({
                    "source": "google-page", "platform": host, "kind": "post", "title": clean(soup.title.string if soup.title else ""),
                    "text": text, "url": url, "verified": "fetched", "query": q,
                })
        time.sleep(1)
    print(f"  pages: {len(out)} paragraphs kept")
    return out


COLLECTORS = {"amazon": collect_amazon, "reddit": collect_reddit, "twitter": collect_twitter, "google": collect_google}


# ---------- import ----------

def read_any(path):
    p = Path(path)
    if p.suffix == ".jsonl":
        return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]
    if p.suffix == ".json":
        d = json.loads(p.read_text(encoding="utf-8"))
        return d if isinstance(d, list) else d.get("rows", [])
    with p.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def normalize_import(r, default_source):
    text = r.get("text") or r.get("quote") or r.get("body") or r.get("review") or r.get("reviewDescription") or ""
    url = r.get("url") or r.get("link") or ""
    kind = r.get("kind") or ("review" if r.get("rating") not in (None, "", "null") or r.get("product") else "post")
    return {
        "source": r.get("source") or default_source, "platform": r.get("platform") or urlparse(url).netloc.replace("www.", ""),
        "kind": kind, "product": r.get("product", ""), "brand": r.get("brand", ""),
        "rating": "" if r.get("rating") in (None, "null") else r.get("rating"),
        "title": clean(r.get("title")), "text": text, "url": url, "author": r.get("author", ""),
        "date": r.get("date", ""), "community": r.get("community", ""),
        "theme": ", ".join(t for t in (r.get("theme"), r.get("complaint_type")) if t), "emotion": r.get("emotion", ""),
        "key_phrases": r.get("key_phrases", ""), "verified": r.get("verified", "imported"),
        "about_product": r.get("about_product", ""),
    }


# ---------- hooks ----------

def split_sentences(text):
    return [s.strip(" \"'") for s in re.split(r"(?<=[.!?])\s+|\n+", text or "") if s.strip()]


def hooks(rows, theme=None, kind=None, n=30, max_words=22, contains=None):
    """Rank exact customer sentences that work as ad hooks: first person, short, emotionally loaded."""
    cands = []
    for r in rows:
        if theme and theme.lower() not in (r.get("theme") or "").lower() and theme.lower() not in (r.get("text") or "").lower():
            continue
        if kind == "reviews" and r.get("kind") != "review":
            continue
        if kind == "forums" and (r.get("kind") == "review" or r.get("about_product") == "yes"):
            continue
        phrases = [p.strip() for p in (r.get("key_phrases") or "").split("|") if p.strip()]
        for s in phrases + split_sentences(r.get("text")):
            words = len(s.split())
            if words < 4 or words > max_words or not FIRST_PERSON.search(s):
                continue
            if contains and contains.lower() not in s.lower():
                continue
            score = pain_score(s) + (3 if s in phrases else 0) + (2 if 6 <= words <= 14 else 0)
            cands.append((score, s, r))
    best, seen = [], set()
    for score, s, r in sorted(cands, key=lambda c: -c[0]):
        k = re.sub(r"\W+", "", s.lower())
        if k in seen:
            continue
        seen.add(k)
        best.append((score, s, r))
        if len(best) >= n:
            break
    return best


# ---------- cli ----------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("collect", help="scrape every source in the config into the sheet")
    c.add_argument("--config", default=str(ROOT / "gut-complex.toml"))
    c.add_argument("--only", help="comma list: amazon,reddit,twitter,google")
    i = sub.add_parser("import", help="add rows from CSV/JSON/JSONL (exports, Claude research, manual copy)")
    i.add_argument("files", nargs="+")
    i.add_argument("--config", default=str(ROOT / "gut-complex.toml"))
    i.add_argument("--source", default="manual")
    h = sub.add_parser("hooks", help="pull hook lines from the sheet")
    h.add_argument("--theme", help="theme tag or any word, e.g. urgency, bloating, relapse, doctor")
    h.add_argument("--kind", choices=["reviews", "forums"], help="only competitor reviews or only forum/social posts")
    h.add_argument("--contains", help="only sentences containing this text")
    h.add_argument("--n", type=int, default=30)
    h.add_argument("--out", help="also write the hooks to this markdown file")
    sub.add_parser("stats", help="row counts by source, kind, theme")
    a = ap.parse_args()

    if a.cmd == "collect":
        cfg = tomllib.loads(Path(a.config).read_text())
        token = os.environ.get("APIFY_TOKEN")
        if not token:
            sys.exit("Set APIFY_TOKEN (https://console.apify.com/account/integrations). See research/README.md.")
        only = a.only.split(",") if a.only else [k for k in COLLECTORS if k in cfg]
        rows = []
        for name in only:
            if name not in cfg:
                print(f"  {name}: not in config, skipped")
                continue
            try:
                rows += COLLECTORS[name](cfg, token)
            except Exception as e:  # one failing source shouldn't lose the others
                print(f"  {name} failed: {e}", file=sys.stderr)
        added, total = merge(rows, cfg)
        print(f"added {added} new rows · sheet now {total} rows → {SHEET_CSV.relative_to(ROOT.parent)}")

    elif a.cmd == "import":
        cfg = tomllib.loads(Path(a.config).read_text()) if Path(a.config).exists() else {}
        rows = [normalize_import(r, a.source) for f in a.files for r in read_any(f)]
        added, total = merge(rows, cfg)
        print(f"added {added} new rows · sheet now {total} rows")

    elif a.cmd == "hooks":
        res = hooks(load_sheet(), a.theme, a.kind, a.n, contains=a.contains)
        lines = [f"{n}. \"{s}\"\n   ({r.get('platform')}{' · ' + r['product'] if r.get('product') else ''}"
                 f"{' · ' + str(r['rating']) + '★' if r.get('rating') else ''}) {r.get('url')}"
                 for n, (_, s, r) in enumerate(res, 1)]
        print("\n".join(lines) or "no matches")
        if a.out:
            Path(a.out).write_text(f"# Hooks{' · ' + a.theme if a.theme else ''}\n\n" + "\n".join(lines) + "\n")

    elif a.cmd == "stats":
        from collections import Counter
        rows = load_sheet()
        print(f"{len(rows)} rows")
        for col in ("source", "kind", "product", "theme", "emotion"):
            vals = [r.get(col) or "—" for r in rows]
            cnt = Counter(t.strip() for v in vals for t in (v.split(",") if col in ("theme", "emotion") else [v]) if t.strip())
            print(f"\n{col}:\n" + "\n".join(f"  {v:>4}  {k}" for k, v in cnt.most_common(15)))


if __name__ == "__main__":
    main()
